from pathlib import Path
from uuid import uuid4

from docx import Document

from app.modules.rag.domain.models import RagDocument
from app.modules.rag.domain.parser import DocumentParser


class DocxParser(DocumentParser):

    def parse(self, file_path: Path) -> list[RagDocument]:
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在：{file_path}")

        if not file_path.is_file():
            raise ValueError(f"传入路径不是文件：{file_path}")

        if file_path.suffix.lower() != ".docx":
            raise ValueError(f"文件格式不是 DOCX：{file_path.suffix}")

        docx_document = Document(str(file_path))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in docx_document.paragraphs
            if paragraph.text.strip()
        ]

        content = "\n".join(paragraphs)

        if not content:
            raise ValueError("Word 文档中没有可用文本")

        document = RagDocument(
            document_id=str(uuid4()),
            content=content,
            source=file_path.name,
            file_type=file_path.suffix.lower(),
            metadata={
                "file_path": str(file_path),
                "paragraph_count": len(paragraphs),
            },
        )

        return [document]
